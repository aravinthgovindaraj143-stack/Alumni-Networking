import os
from datetime import datetime

from django.conf import settings
from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from accounts.models import Alumni


class Command(BaseCommand):

    help = "Import alumni data from Excel"

    def handle(self, *args, **kwargs):

        # Excel file is in the same folder as manage.py
        file_path = os.path.join(
            settings.BASE_DIR,
            "alumni_portal_46_students.xlsx"
        )

        self.stdout.write(
            f"Looking for Excel file at:\n{file_path}"
        )

        # Check whether file exists
        if not os.path.isfile(file_path):

            self.stdout.write(
                self.style.ERROR(
                    "Excel file not found!"
                )
            )

            self.stdout.write(
                "Make sure the Excel file is in the same "
                "folder as manage.py."
            )

            return

        self.stdout.write(
            self.style.SUCCESS(
                "Excel file found!"
            )
        )

        # Open Excel file
        workbook = load_workbook(
            file_path,
            data_only=True
        )

        # Check sheet
        if "Complete Alumni Dataset" not in workbook.sheetnames:

            self.stdout.write(
                self.style.ERROR(
                    "Sheet 'Complete Alumni Dataset' not found!"
                )
            )

            self.stdout.write(
                f"Available sheets: {workbook.sheetnames}"
            )

            return

        worksheet = workbook[
            "Complete Alumni Dataset"
        ]

        rows = list(
            worksheet.iter_rows(
                min_row=2,
                values_only=True
            )
        )

        self.stdout.write(
            f"Found {len(rows)} records in Excel."
        )

        imported = 0
        updated = 0

        for row in rows:

            (
                register_number,
                full_name,
                date_of_birth,
                email,
                phone,
                batch,
                company,
                position,
                feedback,
                kmeans_cluster,
                knn_similar_alumni,
                recommended_alumni,
                sentiment_score,
                sentiment
            ) = row

            # Skip empty rows
            if not register_number or not full_name:
                continue

            # Convert register number to string
            register_number = str(
                register_number
            ).strip()

            full_name = str(
                full_name
            ).strip()

            # Convert Excel datetime to date
            if isinstance(
                date_of_birth,
                datetime
            ):
                date_of_birth = date_of_birth.date()

            # Empty DOB
            if not date_of_birth:
                date_of_birth = None

            # Batch is always 2027
            batch = 2027

            # Create or update record
            alumnus, created = Alumni.objects.update_or_create(

                register_number=register_number,

                defaults={

                    "full_name": full_name,

                    "date_of_birth":
                        date_of_birth,

                    "email": (
                        str(email).strip()
                        if email
                        else ""
                    ),

                    "phone": (
                        str(phone).strip()
                        if phone
                        else ""
                    ),

                    "batch": batch,

                    "company": (
                        str(company).strip()
                        if company
                        else ""
                    ),

                    "position": (
                        str(position).strip()
                        if position
                        else ""
                    ),

                    "feedback": (
                        str(feedback).strip()
                        if feedback
                        else ""
                    ),

                    "kmeans_cluster":
                        kmeans_cluster,

                    "knn_similar_alumni": (
                        str(knn_similar_alumni).strip()
                        if knn_similar_alumni
                        else ""
                    ),

                    "recommended_alumni": (
                        str(recommended_alumni).strip()
                        if recommended_alumni
                        else ""
                    ),

                    "sentiment_score":
                        sentiment_score,

                    "sentiment": (
                        str(sentiment).strip()
                        if sentiment
                        else ""
                    ),
                }
            )

            if created:
                imported += 1
            else:
                updated += 1

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                "========================================"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                "       ALUMNI IMPORT COMPLETED"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                "========================================"
            )
        )

        self.stdout.write(
            f"New records imported : {imported}"
        )

        self.stdout.write(
            f"Existing records updated : {updated}"
        )

        self.stdout.write(
            f"Total records in database : "
            f"{Alumni.objects.count()}"
        )

        self.stdout.write(
            self.style.SUCCESS(
                "Batch set to 2027."
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                "========================================"
            )
        )